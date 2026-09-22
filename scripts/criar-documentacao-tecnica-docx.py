from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "documentacao-tecnica-completa-leste-audio-ia.docx"
BLUE = "0000A3"
YELLOW = "FFD53D"
LIGHT_BLUE = "EAF0FF"
LIGHT_GRAY = "F4F5F7"
GRAY = "5B6470"


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    element = properties.find(qn("w:shd"))
    if element is None:
        element = OxmlElement("w:shd")
        properties.append(element)
    element.set(qn("w:fill"), fill)


def cell_margins(cell, top=100, start=120, bottom=100, end=120):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D9D9D9", size="6"):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_font(run, name="Aptos", size=10.5, color="222222", bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_para(paragraph, before=0, after=6, line=1.12):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_text(paragraph, text, **kwargs):
    run = paragraph.add_run(text)
    set_font(run, **kwargs)
    return run


def add_body(doc, text, after=6):
    paragraph = doc.add_paragraph(style="Normal")
    set_para(paragraph, after=after)
    add_text(paragraph, text)
    return paragraph


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    paragraph = doc.add_paragraph(style=style)
    set_para(paragraph, after=3)
    add_text(paragraph, text)
    return paragraph


def add_number(doc, text, number):
    paragraph = doc.add_paragraph(style="Normal")
    set_para(paragraph, after=3)
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.18)
    add_text(paragraph, f"{number}. ", bold=True)
    add_text(paragraph, text)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    set_para(paragraph, before=12 if level == 1 else 8, after=5, line=1.0)
    run = paragraph.add_run(text)
    set_font(run, name="Aptos Display", size=17 if level == 1 else 12.5, color="000000", bold=True)
    return paragraph


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_margins(cell)
        set_cell_border(cell)
        paragraph = cell.paragraphs[0]
        set_para(paragraph, after=0, line=1.0)
        add_text(paragraph, header, color="FFFFFF", bold=True, size=9.2)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            shade(cell, LIGHT_GRAY if row_index % 2 else "FFFFFF")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)
            set_cell_border(cell)
            paragraph = cell.paragraphs[0]
            set_para(paragraph, after=0, line=1.05)
            add_text(paragraph, str(value), size=9.1)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_code(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.2)
    paragraph.paragraph_format.right_indent = Inches(0.2)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.0
    for line_index, line in enumerate(text.splitlines()):
        if line_index:
            paragraph.add_run().add_break()
        add_text(paragraph, line, name="Cascadia Mono", size=8.5, color="253858")
    return paragraph


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.color.rgb = RGBColor(0, 0, 0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(footer, after=0, line=1.0)
    add_text(footer, "Leste Audio IA | Documentação técnica | ", size=8, color=GRAY)
    add_text(footer, "Uso interno", size=8, color=GRAY, italic=True)


def build_document():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_para(title, before=22, after=4, line=1.0)
    add_text(title, "Leste Audio IA", name="Aptos Display", size=30, color="000000", bold=True)
    subtitle = doc.add_paragraph()
    set_para(subtitle, after=14, line=1.0)
    add_text(subtitle, "Documentação técnica completa do produto", size=15, color=BLUE, bold=True)
    add_body(doc, f"Versão da documentação: {date.today().strftime('%d/%m/%Y')}")
    add_body(doc, "Produto da Leste Valley para transcrição, interpretação, organização e transformação de áudios, documentos e vídeos em conteúdos copiáveis, narráveis e exportáveis.", after=12)

    add_heading(doc, "Visão geral", 1)
    add_body(doc, "O Leste Audio IA é um aplicativo web full stack executado com Next.js e Node.js. O usuário envia áudios, vídeos, PDFs, documentos DOCX, imagens ou texto; o backend processa os arquivos temporariamente; as APIs de IA geram transcrição, resumo, organização, interpretação e conteúdos derivados; e a interface mantém os resultados somente na sessão atual do navegador.")
    add_body(doc, "A aplicação foi estruturada para rodar primeiro em localhost e depois ser publicada em um ambiente Node.js, como a Hostinger. As chaves ficam no servidor por meio de variáveis de ambiente. O pacote de instalação privado pode levar o .env.local atual, mas não deve ser enviado a repositórios ou terceiros.")

    add_heading(doc, "Conclusão operacional", 2)
    add_body(doc, "Para usar o produto em outro computador, instale Node.js 20 ou superior, extraia o pacote, execute o instalador e confira o arquivo .env.local. O instalador baixa as dependências, cria o atalho na Área de Trabalho e inicia o aplicativo em http://localhost:3000.")

    add_heading(doc, "Instalação local", 1)
    add_number(doc, "Instale Node.js 20 ou superior a partir do site oficial do Node.js.", 1)
    add_number(doc, "Extraia leste-audio-ia-distribuicao.zip para uma pasta de trabalho.", 2)
    add_number(doc, "Execute instalar-leste-audio-ia.cmd como usuário comum.", 3)
    add_number(doc, "Preencha as variáveis de ambiente em .env.local antes de usar funções de IA.", 4)
    add_number(doc, "Abra o atalho Leste Audio IA - Iniciar criado na Área de Trabalho.", 5)
    add_code(doc, "npm install\nnpm run dev\n\nAplicação local: http://localhost:3000")
    add_body(doc, "O pacote privado leste-audio-ia-distribuicao-privada.zip já contém o .env.local existente. Ele deve ser tratado como arquivo secreto e usado somente em uma transferência controlada.")

    add_heading(doc, "Variáveis de ambiente", 1)
    add_table(doc, ["Variável", "Finalidade", "Exemplo ou padrão"], [
        ("GEMINI_API_KEY", "Chave do Gemini para áudio, análise multimodal e voz.", "preencher no .env.local"),
        ("GEMINI_MODEL", "Modelo de transcrição ou processamento configurado.", "gemini-3.7-flash"),
        ("DEEPSEEK_API_KEY", "Chave usada para tarefas textuais no backend.", "preencher no .env.local"),
        ("DEEPSEEK_BASE_URL", "Endpoint compatível com OpenAI SDK.", "https://api.deepseek.com"),
        ("DEEPSEEK_MODEL", "Modelo usado para resumo e interpretação.", "deepseek-flash"),
        ("MAX_PARALLEL_TRANSCRIPTIONS", "Quantidade máxima de transcrições paralelas.", "3"),
        ("MAX_FILE_SIZE_MB", "Limite individual para áudio, vídeo e documentos.", "100 ou 125 conforme ambiente"),
        ("TEMP_UPLOAD_DIR", "Pasta temporária para arquivos em processamento.", "./tmp/uploads"),
        ("MAX_VIDEO_UPLOAD_MB", "Limite específico para vídeos.", "125"),
        ("YTDLP_BIN", "Executável opcional para fontes do YouTube.", "yt-dlp"),
    ], widths=[1.7, 3.0, 2.0])
    add_body(doc, "Nunca coloque chaves diretamente em componentes React, rotas públicas, commits ou arquivos enviados ao GitHub. O .gitignore já cobre .env, .env.local, caches, dependências e temporários.")

    add_heading(doc, "Arquitetura", 1)
    add_table(doc, ["Camada", "Responsabilidade", "Principais elementos"], [
        ("Interface", "Upload, fila, cards, resultados, áudio e exportação.", "src/components, src/app/page.tsx"),
        ("Orquestração", "Estado da sessão, fila, progresso e chamadas internas.", "LesteAudioApp.tsx, queue.ts"),
        ("API", "Validação, processamento e integração com provedores.", "src/app/api/*/route.ts"),
        ("IA", "Transcrição, análise textual, intenção, conteúdo e voz.", "src/lib, src/prompts"),
        ("Arquivos", "Temporários, conversão, limpeza e exportação.", "audio.ts, temp-files.ts, export-txt.ts"),
        ("Distribuição", "Instalação e atalho local.", "scripts, instalador e launcher"),
    ], widths=[1.2, 3.4, 2.1])
    add_body(doc, "As API keys nunca chegam ao navegador. A interface chama somente Route Handlers internos do Next.js, que validam entradas e conversam com Gemini, DeepSeek, FFmpeg ou ferramentas locais.")

    add_heading(doc, "Fluxo de áudio", 1)
    add_number(doc, "O usuário seleciona ou arrasta um ou vários arquivos. O navegador apenas mantém os objetos File na sessão até o comando de transcrição.", 1)
    add_number(doc, "A fila valida extensão, MIME type quando possível e tamanho individual. Arquivos inválidos recebem erro próprio e não interrompem os demais.", 2)
    add_number(doc, "Cada arquivo é enviado para POST /api/transcribe. O servidor grava um temporário, tenta a transcrição e sempre limpa os caminhos no finally.", 3)
    add_number(doc, "Se o áudio original falhar, arquivos problemáticos podem ser convertidos para WAV mono 16 kHz com ffmpeg-static ou FFmpeg do sistema.", 4)
    add_number(doc, "Quando a transcrição termina, o aplicativo pode gerar automaticamente resumo, organização e inteligência do áudio. A geração de voz permanece opcional e só ocorre quando o usuário solicita.", 5)
    add_number(doc, "A sessão exibe transcrição individual, cópia individual, resumo, organização, interpretação, voz da Milena e exportações.", 6)
    add_table(doc, ["Status", "Significado"], [
        ("Aguardando ou idle", "Arquivo selecionado, ainda não enviado."),
        ("queued", "Aguardando espaço na fila."),
        ("uploading", "Enviando para o backend."),
        ("converting", "Convertendo com FFmpeg quando necessário."),
        ("transcribing", "Processando o áudio com Gemini."),
        ("done", "Transcrição pronta."),
        ("error", "Falha individual com opção de tentar novamente."),
    ], widths=[1.5, 5.2])

    add_heading(doc, "Cópia e exportação", 1)
    add_body(doc, "Na parte inferior do Painel geral existe o botão Copiar todas as transcrições. Ele reúne somente as transcrições concluídas, identifica cada áudio e não mistura resumo ou interpretação. O botão Copiar tudo da sessão reúne o conteúdo completo disponível.")
    add_table(doc, ["Saída", "Conteúdo"], [
        ("Copiar transcrição", "Somente o texto transcrito de um áudio."),
        ("Copiar todas as transcrições", "Todos os áudios concluídos em uma única cópia."),
        ("TXT", "Transcrições ou organização geral em texto simples, sem Markdown decorativo."),
        ("DOCX", "Transcrições, resumo, organização, interpretação, tarefas, dados-chave, resposta e intenção consolidada."),
        ("MP3", "Áudio gerado pela Milena para trechos e resultados de texto."),
    ], widths=[2.2, 4.5])

    add_heading(doc, "Painel geral e intenção consolidada", 1)
    add_body(doc, "O Painel geral trabalha com todas as transcrições disponíveis. A análise completa executa resumo consolidado, organização, interpretação, tarefas, dados-chave, resposta para WhatsApp e, por último, a intenção consolidada.")
    add_body(doc, "A intenção consolidada não é apenas a intenção de um áudio individual. Ela considera o conjunto inteiro e apresenta assunto, contexto, intenção principal, intenções secundárias, objetivo, ações práticas em ordem de prioridade, requisitos, preferências, restrições, decisões, lacunas, riscos, ambiguidades, sugestões e grau de confiança.")
    add_body(doc, "O resultado diferencia fatos citados, inferências prováveis e sugestões. Uma ação apenas mencionada não é tratada como pedido confirmado sem evidência contextual. A intenção consolidada pode ser copiada, ouvida pela Milena, exportada para TXT/DOCX e usada pelo recurso Transformar em Prompt.")

    add_heading(doc, "Transformar em Prompt", 1)
    add_body(doc, "Depois que há transcrições ou resultados gerais, o usuário pode escolher um dos quatro formatos. O gerador usa o conteúdo interpretado, incluindo a intenção consolidada, para montar uma instrução completa e copiável.")
    add_table(doc, ["Formato", "Uso"], [
        ("Construção de aplicativo", "Produto completo com fluxos, telas, regras e critérios de aceite."),
        ("Criação de app web", "Frontend, backend, APIs, uploads, responsividade, segurança e deploy."),
        ("Criação de landing page", "Promessa, seções, copy, CTAs, SEO e identidade visual."),
        ("Atualização de projeto existente", "Correções, melhorias, compatibilidade e validação sem quebrar o que já funciona."),
    ], widths=[2.2, 4.5])

    add_heading(doc, "PDF, DOCX e audiobook", 1)
    add_body(doc, "O módulo de documentos extrai texto de PDF e DOCX, mostra o conteúdo no painel e permite organizar, resumir, revisar gramática, limpar estrutura e preparar uma versão natural para narração. O audiobook divide o conteúdo em capítulos e permite gerar o áudio por capítulo ou como resultado completo.")
    add_bullet(doc, "PDF: extração de texto e processamento das seções disponíveis.")
    add_bullet(doc, "DOCX: leitura do conteúdo textual e preparação para organização ou narração.")
    add_bullet(doc, "Audiobook: preparação fiel, natural ou adaptada; capítulos; geração progressiva; cancelamento; download por capítulo e download completo.")
    add_bullet(doc, "Privacidade: documentos são tratados como temporários e não formam histórico permanente na primeira versão.")

    add_heading(doc, "Milena Voz", 1)
    add_body(doc, "A área Milena Voz transforma textos, resumos, interpretações, textos de documentos e conteúdos derivados em áudio. O usuário pode escolher estilo e velocidade, ouvir pelo player, pausar, parar e baixar o resultado em MP3 quando o provedor retornar o formato compatível.")
    add_bullet(doc, "A leitura é acionada pelo usuário; a transcrição e a análise não geram áudio automaticamente sem essa escolha.")
    add_bullet(doc, "O áudio fica disponível somente na sessão atual por meio de URL temporária do navegador.")
    add_bullet(doc, "Erros de geração aparecem no bloco correspondente e não apagam a transcrição ou o texto.")

    add_heading(doc, "Vídeos e YouTube", 1)
    add_body(doc, "O aplicativo aceita vídeo MP4 dentro do limite configurado e pode usar a análise multimodal para obter compreensão textual. O módulo Cortes Inteligentes prepara candidatos a cortes de podcast, permite revisão, renderiza MP4 e legenda SRT e expõe saída por endpoint com suporte a Range.")
    add_body(doc, "Para fontes do YouTube, o ambiente precisa ter a ferramenta configurada indicada por YTDLP_BIN. A disponibilidade depende das regras e mudanças do próprio serviço, além de FFmpeg/FFprobe no ambiente de execução.")

    add_heading(doc, "Rotas internas principais", 1)
    add_table(doc, ["Rota", "Uso"], [
        ("GET /api/health", "Confere disponibilidade básica e modelos configurados."),
        ("POST /api/transcribe", "Recebe um arquivo e retorna a transcrição."),
        ("POST /api/summarize", "Resume uma transcrição ou conjunto de textos."),
        ("POST /api/organize", "Organiza e melhora a clareza sem alterar o sentido."),
        ("POST /api/analyze-all", "Interpreta, extrai tarefas, dados, resposta ou intenção geral."),
        ("POST /api/detect-intent", "Detecta a intenção de um texto individual."),
        ("POST /api/speech", "Gera áudio de texto para a Milena."),
        ("POST /api/pdf-extract e pdf-process", "Extrai e processa conteúdo de documentos."),
        ("POST /api/transform-prompt", "Converte o contexto em prompt técnico."),
        ("/api/cuts/jobs/*", "Cria, analisa, renderiza, cancela e serve cortes de vídeo."),
    ], widths=[2.5, 4.2])

    add_heading(doc, "Privacidade e temporários", 1)
    add_body(doc, "A primeira versão não usa banco de dados, login, histórico permanente, localStorage ou IndexedDB para guardar áudio. O backend cria arquivos temporários somente quando necessário e deve removê-los em sucesso e erro. O conteúdo sensível não deve ser impresso no terminal e a transcrição completa não deve ser registrada por padrão.")
    add_body(doc, "O usuário pode clicar em Limpar tudo para remover arquivos, resultados, áudios gerados, resumos, organização, intenção e análise geral da interface. A limpeza da sessão não recupera arquivos que já foram apagados pelo backend.")

    add_heading(doc, "Deploy para Hostinger", 1)
    add_number(doc, "Envie somente o pacote público ou o repositório sem .env.local para o GitHub.", 1)
    add_number(doc, "Conecte o repositório no ambiente Node.js da Hostinger.", 2)
    add_number(doc, "Configure as variáveis de ambiente no painel da Hostinger, sem colar chaves no código.", 3)
    add_number(doc, "Use Node.js 20 ou superior, npm run build e npm run start.", 4)
    add_number(doc, "Confirme que ffmpeg-static funciona no plano escolhido; se não funcionar, instale FFmpeg no VPS ou use um ambiente com suporte adequado.", 5)
    add_code(doc, "npm install\nnpm run build\nnpm run start")
    add_body(doc, "O pacote privado com .env.local não deve ser usado como artefato público de deploy. Para produção, prefira configurar os segredos no painel da Hostinger e rotacionar qualquer chave que tenha sido compartilhada indevidamente.")

    add_heading(doc, "Empacotamento para outro computador", 1)
    add_table(doc, ["Arquivo", "Finalidade"], [
        ("leste-audio-ia-distribuicao.zip", "Pacote recomendado, sem segredos."),
        ("leste-audio-ia-distribuicao-privada.zip", "Pacote controlado com .env.local atual."),
        ("instalar-leste-audio-ia.cmd", "Instala dependências, cria ambiente e gera atalho."),
        ("iniciar-leste-audio-ia.cmd", "Inicia o servidor local e abre o navegador."),
        ("scripts/criar-atalho.ps1", "Cria ou recria o atalho na Área de Trabalho."),
    ], widths=[3.0, 3.7])
    add_code(doc, "npm run package:windows\nnpm run package:windows:private")
    add_body(doc, "O pacote não inclui node_modules de propósito. O instalador executa npm install no computador de destino para obter dependências compatíveis com aquele ambiente.")

    add_heading(doc, "Testes de validação", 1)
    add_table(doc, ["Verificação", "Resultado esperado"], [
        ("npm run typecheck", "TypeScript sem erros."),
        ("npm run lint", "ESLint sem avisos ou erros."),
        ("npm run build", "Build de produção concluído."),
        ("GET /api/health", "Resposta JSON com ok true."),
        ("Upload múltiplo", "Fila continua mesmo com erro individual."),
        ("Cópia inferior", "Todas as transcrições são copiadas de uma vez."),
        ("Análise completa", "Intenção consolidada aparece por último."),
        ("TXT e DOCX", "Conteúdo sem asteriscos e sem títulos Markdown decorativos."),
        ("Milena", "Player permite ouvir, pausar, parar e baixar quando disponível."),
        ("Limpar tudo", "Sessão retorna ao estado inicial."),
    ], widths=[2.5, 4.2])
    add_body(doc, "Checklist manual: abrir localhost; enviar um OGG; transcrever; copiar individual; adicionar vários áudios; copiar todas as transcrições; gerar resumo, organização, interpretação e intenção geral; ouvir pela Milena; baixar MP3, TXT e DOCX; enviar PDF ou DOCX; testar audiobook; limpar tudo; confirmar temporários e .env.local.")

    add_heading(doc, "Limitações conhecidas", 1)
    add_bullet(doc, "Não há integração direta com WhatsApp; os arquivos precisam ser baixados ou encaminhados para o app.")
    add_bullet(doc, "Não há histórico permanente ou banco na primeira versão.")
    add_bullet(doc, "A qualidade depende do áudio, ruído, sobreposição de vozes e distância do microfone.")
    add_bullet(doc, "Arquivos muito grandes podem exigir divisão futura em partes; o limite deve ser ajustado conforme o ambiente.")
    add_bullet(doc, "A disponibilidade de modelos e formatos de voz depende da conta, do endpoint e das políticas do provedor.")
    add_bullet(doc, "Cortes de YouTube dependem de yt-dlp, FFmpeg, FFprobe e das condições do vídeo.")

    add_heading(doc, "Estrutura de manutenção", 1)
    add_body(doc, "Ao alterar uma integração, preserve o contrato JSON entre componente e Route Handler, atualize os tipos em src/types/audio.ts, mantenha mensagens de erro individuais e execute typecheck, lint e build. Alterações em DOCX, PDF, áudio ou vídeo devem incluir um teste manual do fluxo correspondente.")
    add_body(doc, "Antes de publicar, confirme que nenhuma chave aparece no Git, que tmp e caches estão ignorados, que o domínio aponta para o ambiente correto e que o servidor possui FFmpeg quando a funcionalidade precisar de conversão.")

    doc.core_properties.title = "Leste Audio IA Documentação Técnica Completa"
    doc.core_properties.subject = "Arquitetura, instalação, uso, testes e operação"
    doc.core_properties.author = "Leste Valley"
    doc.core_properties.keywords = "Leste Audio IA, transcrição, áudio, Milena, Next.js, documentação"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
