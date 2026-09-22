from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "documentacao-tecnica-leste-audio-ia.docx"

BLUE = "0000A3"
GOLD = "FFD53D"
LIGHT_BLUE = "EAF0FF"
LIGHT_GRAY = "F4F6F8"
MID_GRAY = "D9D9D9"
TEXT = "1F2937"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = MID_GRAY) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 100, start: int = 110, bottom: int = 100, end: int = 110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_font(run, name: str = "Aptos", size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    """Keep each compact table row together when LibreOffice paginates it."""
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size in (("Title", 27), ("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 10.8)):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(15 if style_name == "Heading 1" else 10)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in [style.name for style in doc.styles]:
        style = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Consolas"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        style.font.size = Pt(8.6)
        style.font.color.rgb = RGBColor.from_string(TEXT)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.0

    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header_p.add_run("LESTE VALLEY  |  LESTE AUDIO IA  |  DOCUMENTAÇÃO TÉCNICA")
    set_font(run, size=8, bold=True, color="4B5563")

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run("Leste Audio IA  |  Página ")
    set_font(run, size=8, color="4B5563")
    add_page_field(footer_p)


def add_paragraph(doc: Document, text: str = "", bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_lead:
        run = p.add_run(bold_lead)
        set_font(run, bold=True)
    if text:
        run = p.add_run(text)
        set_font(run)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(item)
        set_font(run)


def add_numbered(doc: Document, items: list[str]) -> None:
    # Word's built-in list style may continue the numbering from the table of
    # contents. Emit the number explicitly so every documented flow starts at 1.
    for position, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.space_after = Pt(2)
        number = p.add_run(f"{position}. ")
        set_font(number)
        run = p.add_run(item)
        set_font(run)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        set_cell_shading(cell, BLUE)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        run = cell.paragraphs[0].add_run(header)
        set_font(run, size=8.7, bold=True, color=WHITE)
        if widths:
            cell.width = Inches(widths[index])

    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cell = cells[index]
            cell.text = ""
            set_cell_border(cell)
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, LIGHT_BLUE)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            set_font(run, size=8.5, color=TEXT)
            if widths:
                cell.width = Inches(widths[index])
        prevent_row_split(table.rows[-1])

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.text = ""
    set_cell_shading(cell, LIGHT_GRAY)
    set_cell_border(cell)
    set_cell_margins(cell, top=120, start=150, bottom=120, end=150)
    p = cell.paragraphs[0]
    p.style = doc.styles["Code Block"]
    for line_number, line in enumerate(text.strip().splitlines()):
        if line_number:
            p.add_run().add_break()
        run = p.add_run(line)
        set_font(run, name="Consolas", size=8.5, color=TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)


def page_break(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(105)
    p.paragraph_format.space_after = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("LESTE VALLEY")
    set_font(run, size=12, bold=True, color=BLUE)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Leste Audio IA")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(26)
    run = subtitle.add_run("Documentação técnica completa do sistema")
    set_font(run, size=14, color=TEXT)

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = [
        ("Produto", "Leste Audio IA"),
        ("Escopo", "Aplicação web para transcrição, entendimento e voz"),
        ("Versão documental", "1.0"),
        ("Atualização", date.today().strftime("%d/%m/%Y")),
    ]
    for idx, (label, value) in enumerate(rows):
        for cell in table.rows[idx].cells:
            set_cell_border(cell)
            set_cell_margins(cell, top=130, start=140, bottom=130, end=140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(table.rows[idx].cells[0], BLUE)
        table.rows[idx].cells[0].text = ""
        table.rows[idx].cells[1].text = ""
        label_run = table.rows[idx].cells[0].paragraphs[0].add_run(label)
        set_font(label_run, size=9, bold=True, color=WHITE)
        value_run = table.rows[idx].cells[1].paragraphs[0].add_run(value)
        set_font(value_run, size=9, color=TEXT)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(38)
    run = p.add_run("Documento de referência para operação, manutenção e evolução técnica.")
    set_font(run, size=10, color="4B5563")

    page_break(doc)


def build_document() -> Document:
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    heading(doc, "Como usar este documento")
    add_paragraph(
        doc,
        "Este documento descreve como o Leste Audio IA é organizado, executado e mantido. Ele é destinado a pessoas que operam o sistema, desenvolvedores responsáveis por correções e evolução, e responsáveis pelo ambiente de deploy.",
    )
    add_paragraph(
        doc,
        "O sistema roda como uma aplicação Next.js com backend Node.js. Arquivos de áudio, vídeo, PDF e imagem são enviados sob demanda, processados em memória ou em diretório temporário e removidos após cada operação. Nenhuma chave de API é enviada ao navegador.",
    )

    heading(doc, "Sumário")
    add_numbered(doc, [
        "Visão geral e capacidades",
        "Arquitetura do sistema",
        "Fluxos de processamento",
        "Modelos e integrações de inteligência artificial",
        "Rotas internas e contratos de API",
        "Configuração por variáveis de ambiente",
        "Arquivos temporários privacidade e segurança",
        "Voz da Milena e geração de MP3",
        "Painel geral PDF imagem texto e YouTube",
        "Operação local validação e diagnóstico",
        "Deploy GitHub e Hostinger",
        "Limitações e evolução recomendada",
    ])

    heading(doc, "Visão geral e capacidades")
    add_paragraph(
        doc,
        "Leste Audio IA é um web app para receber áudios e vídeos enviados manualmente pelo usuário, produzir transcrições, organizar conteúdo, extrair tarefas e dados, interpretar conjuntos de transcrições e gerar novos materiais. A identidade e a interface pertencem ao Leste Valley.",
    )
    add_table(doc, ["Área", "O que entrega"], [
        ["Upload de áudio e vídeo", "Recebe OGG, OPUS, M4A, MP3, WAV, WEBM, AAC, FLAC e MP4, com limite configurável por arquivo."],
        ["Transcrição", "Transcreve um arquivo por requisição, processa vários itens em fila e continua o lote quando um item falha."],
        ["Painel geral", "Gera resumo consolidado, organização por temas, interpretação, tarefas, dados-chave, decisões e resposta para WhatsApp."],
        ["Entendimento avançado", "Analisa o arquivo original para identificar assunto, contexto, intenção, temas, insights, riscos, oportunidades, narrativa e conteúdos possíveis."],
        ["Geração de conteúdo", "Cria artigo, copy, post, roteiro, storytelling, resposta para WhatsApp e plano de audiolivro a partir da análise e transcrição."],
        ["Voz da Milena", "Lê textos, resumos, análises, PDFs, conteúdo de imagem e materiais gerados; permite ouvir, pausar, parar e baixar em MP3."],
        ["PDF e imagem", "Extrai texto de PDF e imagem, permite organizar, resumir, corrigir, narrar e preparar audiolivro."],
        ["YouTube", "Obtém informações e analisa conteúdo de vídeos do YouTube por rotas específicas do backend."],
        ["Exportação", "Baixa sessões e conteúdos organizados em TXT e DOCX de texto puro, sem marcadores Markdown."],
    ], [1.55, 5.85])

    heading(doc, "Arquitetura do sistema")
    add_paragraph(doc, "A arquitetura separa a interface cliente do processamento sensível no servidor. O frontend guarda apenas o estado da sessão atual do navegador. As rotas internas fazem a validação, chamam provedores externos e respondem em JSON ou mídia binária.")
    add_table(doc, ["Camada", "Tecnologia", "Responsabilidade"], [
        ["Interface", "Next.js App Router, React, TypeScript e Tailwind CSS", "Upload, fila visual, cards, cópia, players, exportação e estados de erro."],
        ["Backend", "Route Handlers do Next.js em runtime Node.js", "Validação, multipart, arquivos temporários, integração de APIs e tratamento de falhas."],
        ["Transcrição", "Gemini Interactions API", "Gemini 3.5 Transcribe para reconhecimento de fala em português brasileiro."],
        ["Entendimento multimodal", "Gemini Interactions API", "Gemini 3.8 Flash para analisar áudio ou vídeo original e responder em JSON estruturado."],
        ["Texto", "Gemini ou DeepSeek opcional", "Resumo, organização, interpretação, prompts, correções e conteúdo derivado."],
        ["Voz", "Gemini TTS e conversão MP3", "Geração de voz da Milena e entrega em audio/mpeg."],
        ["Conversão", "ffmpeg-static com fallback do sistema", "Converte formatos problemáticos para WAV mono a 16 kHz quando necessário."],
    ], [1.4, 1.8, 4.2])

    heading(doc, "Estrutura de pastas relevante")
    add_code(doc, "src/\n  app/\n    api/\n      transcribe/\n      audio-analyze/\n      audio-content/\n      speech/\n      pdf-extract/\n      pdf-process/\n      image-extract/\n      youtube-analyze/\n  components/\n    LesteAudioApp.tsx\n    AudioList.tsx\n    AudioCard.tsx\n    AudioIntelligencePanel.tsx\n    ResultPanel.tsx\n    MilenaVozPanel.tsx\n  lib/\n    env.ts\n    gemini.ts\n    audio-intelligence.ts\n    gemini-tts.ts\n    audio.ts\n    temp-files.ts\n    text-ai.ts\n  prompts/\n  types/")
    add_paragraph(doc, "O componente LesteAudioApp concentra o estado da sessão e conecta os painéis. AudioCard mostra um arquivo individual. AudioIntelligencePanel exibe a análise multimodal e os materiais derivados. As bibliotecas em src/lib não podem ser importadas no navegador quando manipulam chaves ou arquivos.")

    heading(doc, "Fluxos de processamento")
    heading(doc, "Transcrição de áudio ou vídeo", 2)
    add_numbered(doc, [
        "O usuário seleciona ou arrasta arquivos. O navegador valida extensão e tamanho antes de enviar.",
        "O arquivo permanece no navegador até o clique em Transcrever áudios.",
        "A fila processa no máximo MAX_PARALLEL_TRANSCRIPTIONS itens em paralelo. Um erro não interrompe os demais.",
        "A rota POST /api/transcribe recebe somente um arquivo multipart por vez e salva uma cópia temporária com nome sanitizado e único.",
        "O backend envia o arquivo ao Gemini Files API e espera o estado ACTIVE antes do uso.",
        "A transcrição é solicitada pelo Gemini 3.5 Transcribe usando Interactions API, com language_codes pt-BR, modo smart e store false.",
        "Se o modelo dedicado não estiver disponível, a integração preserva o caminho legado com GEMINI_MODEL. Em formatos problemáticos, a rota pode converter para WAV e repetir.",
        "O arquivo local, arquivos convertidos e arquivo remoto do Gemini são removidos no finally, inclusive quando ocorre erro.",
    ])

    heading(doc, "Análise avançada do arquivo original", 2)
    add_numbered(doc, [
        "Após haver uma transcrição, o usuário seleciona Analisar áudio no card individual.",
        "O arquivo original é enviado temporariamente para POST /api/audio-analyze.",
        "Gemini 3.8 Flash recebe um prompt textual e o conteúdo de áudio ou vídeo. A rota solicita resposta estruturada e usa store false.",
        "O resultado é validado com Zod antes de chegar ao frontend.",
        "O painel apresenta assunto, contexto, intenção, confiança, resumo, temas, pontos principais, insights, tarefas, decisões, entidades, riscos, oportunidades, dúvidas, estrutura narrativa e conteúdos sugeridos.",
    ])

    heading(doc, "Criação de conteúdo a partir de áudio", 2)
    add_paragraph(doc, "POST /api/audio-content recebe a transcrição e o objeto de análise já validado. O prompt obriga fidelidade à fonte e entrega texto simples. A geração pode usar Gemini por padrão ou DeepSeek quando TEXT_AI_PROVIDER=deepseek.")
    add_table(doc, ["Formato", "Uso principal"], [
        ["article", "Artigo com título, abertura, desenvolvimento e conclusão prática."],
        ["copy", "Texto persuasivo sem promessas ou fatos não sustentados."],
        ["post", "Conteúdo social com gancho, desenvolvimento e chamada para ação."],
        ["script", "Roteiro narrável com abertura, blocos e encerramento."],
        ["storytelling", "Narrativa baseada nos fatos, com lacunas mantidas como dúvida."],
        ["whatsapp", "Mensagem objetiva e adequada ao contexto identificado."],
        ["audiobookOutline", "Plano de audiolivro com premissa, capítulos e orientação de narração."],
    ], [1.7, 5.7])

    heading(doc, "Modelos e integrações de inteligência artificial")
    add_table(doc, ["Variável", "Modelo padrão", "Uso"], [
        ["GEMINI_TRANSCRIBE_MODEL", "gemini-3.5-transcribe", "Transcrição dedicada, com reconhecimento de idioma e modo smart."],
        ["GEMINI_AUDIO_UNDERSTANDING_MODEL", "gemini-3.8-flash", "Compreensão multimodal de áudio e vídeo e resposta estruturada."],
        ["GEMINI_TEXT_MODEL", "gemini-3.8-flash", "Resumo, organização, interpretação, prompts, PDF, imagem e conteúdos derivados."],
        ["GEMINI_MODEL", "gemini-3.7-flash", "Contingência para a integração legada de transcrição."],
        ["GEMINI_TTS_MODEL", "gemini-3.1-flash-tts-preview", "Leitura regular da Milena."],
        ["GEMINI_TTS_AUDIOBOOK_MODEL", "gemini-2.5-pro-preview-tts", "Audiobooks e narrações longas."],
        ["GEMINI_LIVE_MODEL", "gemini-3.1-flash-live-preview", "Reservado para uma futura interface de conversação em tempo real."],
        ["DEEPSEEK_MODEL", "deepseek-v4-pro", "Alternativa opcional para tarefas somente textuais."],
    ], [2.1, 2.25, 3.05])
    add_paragraph(doc, "A transcrição, a análise multimodal e o TTS utilizam Gemini. DeepSeek não recebe áudio nem vídeo; quando ativado, recebe apenas texto já extraído ou transcrito. Nenhum raciocínio oculto é solicitado ou exibido ao usuário.")

    heading(doc, "Rotas internas e contratos de API")
    add_table(doc, ["Rota", "Método", "Entrada", "Saída principal"], [
        ["/api/health", "GET", "Sem corpo", "Status, modelos ativos e configuração pública."],
        ["/api/transcribe", "POST", "multipart file", "ok, transcription, meta com nome, conversão e modelo."],
        ["/api/audio-analyze", "POST", "multipart file", "Análise multimodal estruturada e modelo usado."],
        ["/api/audio-content", "POST", "kind, transcription, analysis", "Texto derivado pronto para copiar."],
        ["/api/summarize", "POST", "text, mode", "Resumo individual ou consolidado."],
        ["/api/organize", "POST", "text, mode", "Texto reorganizado individual ou consolidado."],
        ["/api/analyze-all", "POST", "items, mode", "Interpretação, tarefas, dados-chave ou resposta."],
        ["/api/speech", "POST", "text, format, estilo", "MP3 da voz da Milena."],
        ["/api/pdf-extract", "POST", "multipart PDF", "Texto extraído do PDF."],
        ["/api/pdf-process", "POST", "text, mode", "Resumo, organização, gramática ou limpeza."],
        ["/api/image-extract", "POST", "multipart image", "Texto identificado na imagem."],
        ["/api/youtube-info", "POST", "URL", "Metadados de vídeo YouTube."],
        ["/api/youtube-analyze", "POST", "URL e opções", "Transcrição e análise de YouTube."],
        ["/api/generate-prompt", "POST", "Textos e formato", "Prompt técnico estruturado."],
        ["/api/transform-prompt", "POST", "Análise e formato", "Prompt final para app, web app, landing page ou atualização."],
        ["/api/detect-intent", "POST", "text", "JSON de intenção e instrução convertida."],
        ["/api/analyze-context", "POST", "text e modo", "Análise unificada de contexto."],
        ["/api/text-correct", "POST", "text e modo", "Texto corrigido ou reescrito."],
        ["/api/prepare-narration", "POST", "text e modo", "Texto preparado para narração."],
        ["/api/export-txt", "POST", "Dados de sessão", "Arquivo TXT quando usado pelo cliente."],
    ], [1.55, 0.55, 2.35, 2.95])

    heading(doc, "Detalhe dos contratos principais", 2)
    add_code(doc, "POST /api/transcribe\nContent-Type: multipart/form-data\nCampo obrigatório: file\n\nResposta de sucesso\n{\n  \"ok\": true,\n  \"transcription\": \"...\",\n  \"meta\": {\n    \"originalFileName\": \"audio.ogg\",\n    \"converted\": false,\n    \"model\": \"gemini-3.5-transcribe\"\n  }\n}")
    add_code(doc, "POST /api/audio-content\nContent-Type: application/json\n{\n  \"kind\": \"audiobookOutline\",\n  \"transcription\": \"...\",\n  \"analysis\": { \"subject\": \"...\", \"...\": \"estrutura completa\" }\n}\n\nResposta\n{ \"ok\": true, \"result\": \"...\", \"model\": \"gemini-3.8-flash\" }")
    add_paragraph(doc, "Em erro, as rotas respondem com ok false e error. A interface mostra a mensagem no bloco afetado sem derrubar o restante da sessão.")

    heading(doc, "Configuração por variáveis de ambiente")
    add_code(doc, "GEMINI_API_KEY=\nGEMINI_MODEL=gemini-3.7-flash\nGEMINI_TRANSCRIBE_MODEL=gemini-3.5-transcribe\nGEMINI_AUDIO_UNDERSTANDING_MODEL=gemini-3.8-flash\nGEMINI_TEXT_MODEL=gemini-3.8-flash\nGEMINI_TTS_MODEL=gemini-3.1-flash-tts-preview\nGEMINI_TTS_AUDIOBOOK_MODEL=gemini-2.5-pro-preview-tts\nGEMINI_LIVE_MODEL=gemini-3.1-flash-live-preview\nGEMINI_TTS_VOICE=Kore\nTEXT_AI_PROVIDER=gemini\n\nDEEPSEEK_API_KEY=\nDEEPSEEK_BASE_URL=https://api.deepseek.com\nDEEPSEEK_MODEL=deepseek-v4-pro\n\nMAX_PARALLEL_TRANSCRIPTIONS=3\nMAX_FILE_SIZE_MB=125\nTEMP_UPLOAD_DIR=./tmp/uploads\nAPP_NAME=Leste Audio IA")
    add_table(doc, ["Grupo", "Comportamento"], [
        ["Segredos", "GEMINI_API_KEY e DEEPSEEK_API_KEY ficam apenas em .env.local ou no painel do host. Nunca devem entrar no repositório ou frontend."],
        ["Texto", "TEXT_AI_PROVIDER=gemini usa Gemini. TEXT_AI_PROVIDER=deepseek transfere apenas tarefas textuais para DeepSeek."],
        ["Limites", "MAX_FILE_SIZE_MB controla a validação no frontend e no backend. O valor atual é 125 MB por arquivo."],
        ["Paralelismo", "MAX_PARALLEL_TRANSCRIPTIONS limita a fila da sessão para preservar estabilidade e custo."],
        ["Temporários", "TEMP_UPLOAD_DIR deve apontar para uma pasta gravável e descartável. O padrão é ./tmp/uploads."],
    ], [1.4, 5.9])

    heading(doc, "Arquivos temporários privacidade e segurança")
    add_bullets(doc, [
        "Não existe banco de dados, login, histórico permanente, localStorage ou IndexedDB para armazenar áudio.",
        "O arquivo só sai do navegador após ação explícita do usuário.",
        "saveUploadedFileToTemp cria nome com nanoid e sanitiza o nome original para impedir path traversal.",
        "A limpeza local usa try catch finally e ignora arquivos já removidos.",
        "A integração Gemini espera o arquivo remoto ficar ACTIVE e o remove no finally depois do processamento.",
        "As operações pela Interactions API usam store false para não preservar estado de interação no provedor.",
        "O backend não registra conteúdo integral de áudio, transcrição ou chave no terminal por padrão.",
        "A API key ausente retorna a mensagem Configure a chave da API no arquivo .env.local sem quebrar a tela inteira.",
    ])
    add_paragraph(doc, "A aplicação não substitui uma política corporativa de retenção ou consentimento. Antes de usar material de terceiros, a operação deve verificar autorização, confidencialidade, LGPD e políticas internas do Leste Valley.")

    heading(doc, "Processamento de formatos e FFmpeg")
    add_table(doc, ["Item", "Regra"], [
        ["Extensões", "ogg, opus, m4a, mp3, wav, webm, aac, flac e mp4."],
        ["MIME", "É validado quando disponível. Valores genéricos application/octet-stream são aceitos quando a extensão é permitida."],
        ["Conversão", "OGG, OPUS, WEBM, M4A, AAC e MP4 podem ser convertidos em caso de incompatibilidade."],
        ["Saída de conversão", "WAV mono a 16 kHz: ffmpeg -y -i input -ac 1 -ar 16000 output.wav."],
        ["Ordem", "ffmpeg-static é tentado primeiro; se falhar, o app busca ffmpeg instalado no sistema."],
        ["Falha", "A rota devolve erro claro. Arquivos originais e convertidos continuam cobertos pela limpeza no finally."],
    ], [1.65, 5.65])

    heading(doc, "Voz da Milena e geração de MP3")
    add_paragraph(doc, "A rota POST /api/speech recebe texto, título opcional, formato, velocidade, estilo e dicionário de pronúncia. Ela divide textos extensos em blocos, envia cada bloco ao TTS, junta o resultado quando possível e entrega audio/mpeg. Se o provedor devolver WAV, o backend codifica MP3 em memória.")
    add_table(doc, ["Estilo", "Uso recomendado"], [
        ["natural", "Textos curtos e leitura regular."],
        ["didatica", "Resumos, instruções e materiais explicativos."],
        ["audiolivro", "Livros, capítulos e narrativas longas. Seleciona GEMINI_TTS_AUDIOBOOK_MODEL."],
        ["podcast e roteiro", "Conteúdos com ritmo conversacional ou estrutura de apresentação."],
        ["formal e institucional", "Comunicados, documentos e materiais corporativos."],
    ], [1.85, 5.45])
    add_paragraph(doc, "Os players do painel exibem controles nativos. Ao selecionar Pausar / parar, o componente pausa o elemento de áudio e retorna ao início; não apenas remove um indicador visual. O download utiliza o Blob gerado na sessão e salva um arquivo MP3 no dispositivo do usuário.")

    heading(doc, "Painel geral PDF imagem texto e YouTube")
    add_table(doc, ["Fonte", "Fluxo"], [
        ["Painel geral", "Une as transcrições disponíveis. Pode resumir, organizar, interpretar, extrair tarefas, mapear dados e gerar resposta para WhatsApp."],
        ["Transformar em Prompt", "Usa transcrições e análise para gerar instruções para aplicativo, app web, landing page ou atualização de projeto existente."],
        ["PDF", "Extrai texto, oferece resumo, organização, correção gramatical e limpeza; pode dividir conteúdo e gerar audiolivro por capítulos."],
        ["Imagem", "Envia a imagem ao backend, extrai ou descreve texto e permite ouvir e baixar a leitura."],
        ["Texto livre", "Permite colar texto, corrigir, analisar, gerar conteúdo e converter em voz."],
        ["YouTube", "Consulta informações e processa conteúdo por rotas específicas, mantendo o mesmo padrão de análise e exportação."],
    ], [1.55, 5.75])

    heading(doc, "Operação local")
    add_paragraph(doc, "O projeto requer Node.js 20 ou superior. Node 20 LTS é a versão recomendada para Next.js 14. O ambiente local atual também funciona com Node mais recente, mas essa combinação não é a referência de compatibilidade para produção.")
    add_code(doc, "npm install\nnpm run dev\n\nAbrir\nhttp://localhost:3000")
    add_paragraph(doc, "No Windows, o atalho iniciar-leste-audio-ia.cmd chama iniciar-leste-audio-ia.ps1. O iniciador verifica GET /api/health, carrega a página inicial e confirma arquivos estáticos. Se o servidor existente estiver inconsistente, ele encerra somente processos pertencentes a este projeto, limpa .next-dev e sobe outra instância.")
    add_code(doc, "iniciar-leste-audio-ia.cmd\niniciar-leste-audio-ia.cmd -Reiniciar\niniciar-leste-audio-ia.cmd -Limpar\niniciar-leste-audio-ia.cmd -NoOpen")

    heading(doc, "Validação e testes")
    add_code(doc, "npm run typecheck\nnpm run lint\nnpm run build")
    add_table(doc, ["Teste", "Resultado esperado"], [
        ["Health check", "GET /api/health retorna ok true e os modelos configurados sem revelar chaves."],
        ["Transcrição", "Enviar arquivo permitido retorna transcription e meta; falha em um item não interrompe o lote."],
        ["Análise avançada", "POST /api/audio-analyze retorna campos estruturados e remove temporários."],
        ["Conteúdo derivado", "POST /api/audio-content retorna texto simples no formato escolhido."],
        ["Voz", "POST /api/speech entrega Content-Type audio/mpeg e arquivo com tamanho maior que zero."],
        ["Privacidade", "Após as operações, tmp/uploads não deve conter arquivos temporários criados pelo teste."],
        ["Interface", "Validar no desktop e mobile upload, estados, cópia, player, pausa, parada, download e limpar tudo."],
    ], [1.75, 5.55])
    add_paragraph(doc, "Validação realizada na configuração atual: health check retornou Gemini 3.5 Transcribe, Gemini 3.8 Flash para análise e texto, Gemini 3.1 Flash TTS e Gemini 2.5 Pro TTS para audiolivro. Foi exercitado um MP3 gerado localmente em transcrição, análise, conteúdo de audiolivro e voz. As verificações npm run typecheck, npm run lint e npm run build passaram.")

    heading(doc, "Diagnóstico de falhas")
    add_table(doc, ["Sintoma", "Causa provável", "Ação"], [
        ["Chave ausente", "GEMINI_API_KEY ou DEEPSEEK_API_KEY não configurada.", "Criar ou corrigir .env.local e reiniciar o servidor."],
        ["Permission denied 403", "Chave sem acesso ao projeto Gemini ou com restrição incompatível.", "Validar a chave no Google AI Studio e remover restrições incompatíveis com o servidor."],
        ["Arquivo não ACTIVE", "Files API ainda processando ou falhou no upload.", "Aguardar a rotina de polling; revisar formato e chave se houver estado FAILED."],
        ["Arquivo acima do limite", "MAX_FILE_SIZE_MB menor que o arquivo.", "Ajustar variável e reiniciar; verificar memória e limites do host."],
        ["EEXIST em temporários", "TEMP_UPLOAD_DIR aponta para arquivo, não pasta, ou está inválido.", "Usar uma pasta gravável, por exemplo ./tmp/uploads."],
        ["Falha FFmpeg", "ffmpeg-static indisponível ou host sem binário compatível.", "Verificar o pacote no deploy; instalar FFmpeg no VPS quando necessário."],
        ["Áudio continua tocando", "Player não recebeu estado de parada ou navegador manteve reprodução.", "Usar Pausar / parar; o painel atual pausa e zera o player explicitamente."],
        ["Tela em branco local", "Cache de desenvolvimento ou chunks antigos.", "Executar iniciar-leste-audio-ia.cmd -Reiniciar -Limpar."],
    ], [1.45, 2.55, 3.3])

    heading(doc, "Deploy GitHub e Hostinger")
    add_numbered(doc, [
        "Confirme que .env.local, tmp, node_modules, .next, dist e out estão no .gitignore.",
        "Envie o código para o repositório GitHub sem chaves ou arquivos temporários.",
        "No ambiente Node.js da Hostinger, configure as mesmas variáveis do arquivo .env.example pelo painel de ambiente.",
        "Defina Node.js 20 ou superior, execute npm install e use npm run build como comando de build.",
        "Use npm run start como comando de produção. O server.js inicia Next em modo produção e escuta PORT fornecida pelo host.",
        "Confirme permissão de escrita para TEMP_UPLOAD_DIR e descarte a pasta tmp em cada implantação, quando apropriado.",
        "Verifique o funcionamento de ffmpeg-static. Se o binário não for compatível, instale FFmpeg no VPS ou ajuste o ambiente." ,
        "Após publicar, valide /api/health, transcrição, análise multimodal, voz MP3 e uma exportação antes de liberar para uso." ,
    ])
    add_paragraph(doc, "O domínio planejado é https://lesteaudio.space. O deploy não deve colocar chaves no repositório, em arquivos públicos, em variáveis NEXT_PUBLIC ou em código JavaScript enviado ao navegador.")

    heading(doc, "Limitações atuais")
    add_bullets(doc, [
        "Não existe integração direta com WhatsApp. O usuário precisa exportar, baixar ou encaminhar o arquivo antes do upload.",
        "Não existe login, banco de dados ou histórico permanente.",
        "Não há divisão automática de arquivos grandes. O limite é configurável, mas arquivos longos dependem de memória, timeout e limites dos provedores.",
        "A qualidade da transcrição depende do áudio. Ruído, sobreposição de vozes e fala distante podem gerar trechos menos precisos.",
        "A análise multimodal produz apoio editorial e operacional. Ela não substitui revisão humana para decisões jurídicas, financeiras, médicas ou estratégicas.",
        "Live API está apenas preparada por configuração. Conversação de microfone em tempo real ainda não está exposta na interface.",
    ])

    heading(doc, "Evolução recomendada")
    add_bullets(doc, [
        "Adicionar controles de transcrição verbatim, diarização e timestamps quando houver necessidade de auditoria de fala.",
        "Criar fila persistente e jobs em segundo plano para audiobooks longos e uploads muito grandes.",
        "Adicionar medições de custo, duração e tokens por operação sem registrar conteúdo sensível.",
        "Incluir testes automatizados de rotas com mocks para Gemini, DeepSeek e FFmpeg.",
        "Implementar monitoramento de erros do backend e health check externo no ambiente de produção.",
        "Criar interface Live usando token efêmero no backend para conversação de voz com interrupção dinâmica.",
    ])

    heading(doc, "Referências técnicas")
    add_bullets(doc, [
        "Gemini API Audio transcription: https://ai.google.dev/gemini-api/docs/transcribe",
        "Gemini API Audio understanding: https://ai.google.dev/gemini-api/docs/audio",
        "Gemini API Speech generation: https://ai.google.dev/gemini-api/docs/speech-generation",
        "Gemini API Interactions: https://ai.google.dev/gemini-api/docs/interactions-overview",
        "Gemini API Files: https://ai.google.dev/gemini-api/docs/files",
        "Next.js Route Handlers: https://nextjs.org/docs/app/building-your-application/routing/route-handlers",
    ])

    return doc


if __name__ == "__main__":
    document = build_document()
    document.core_properties.title = "Leste Audio IA Documentação Técnica"
    document.core_properties.subject = "Arquitetura, operação e manutenção do Leste Audio IA"
    document.core_properties.author = "Leste Valley"
    document.core_properties.keywords = "Leste Audio IA, Next.js, Gemini, áudio, voz, documentação técnica"
    document.save(OUTPUT)
    print(OUTPUT)
